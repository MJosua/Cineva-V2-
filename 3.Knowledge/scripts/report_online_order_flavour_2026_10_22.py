from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

target = Path(__file__).resolve().parents[1] / 'reports/online_order_flavour_2026_10_22.xlsx'
target.parent.mkdir(parents=True, exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'Audit'
rows = [
    ['Online Order Kantor', '22 Oktober 2026', 'Patch lokal; belum deploy/browser'],
    ['Temuan', 'Bukti', 'Batas validasi'],
    ['Reset bentrok diterapkan', 'Handler membandingkan String(SKU)', 'Slot dependen dan metadata dibersihkan'],
    ['Filter tetap anti-duplikat', 'Filter menolak SKU slot sebelumnya', 'Opsi aktif dipertahankan'],
    ['Tes handler aktual lulus', 'Reproduksi gagal sebelum, lulus setelah patch', 'Belum browser authenticated'],
    ['Regresi Container Groups', '9 tes existing lulus', 'Tidak membuktikan submit backend'],
    ['Build Vite lulus', '1840 modules; 21.29 detik', 'Warning pdfjs eval dan ukuran bundle'],
    ['Filter anti-duplikat disengaja', 'FlavourNonBulk_Component.jsx:731-752', 'includes sensitif tipe SKU'],
    ['Ganti Flavour 1: SKU kedua tetap', 'FlavourNonBulk_Component.jsx:533-536', 'Non-pallet hanya quantity menjadi nol'],
    ['Jalur pallet reset SKU ke 0', 'FlavourNonBulk_Component.jsx:550-551', 'Belum konsisten dengan placeholder -1'],
    ['Choose Flavour menjalankan reset', 'FlavourNonBulk_Component.jsx:430-454', 'Sesuai workaround pengguna'],
    ['Nilai kosong tidak konsisten', 'FlavourNonBulk_Component.jsx:554-584', '0, string kosong, dan -1'],
    ['Belum diperiksa', 'Browser production, payload, DB', 'Penyebab transaksi belum terbukti'],
    ['Source', 'D:/Kantor/1.Frontend/E-Order/src/features/orders/components/AddMoreContainer/FlavourNonBulk_Component.jsx', 'Kantor read-only'],
]
for row in rows:
    ws.append(row)
for col in 'ABC':
    ws.column_dimensions[col].width = 52
for row in ws:
    ws.row_dimensions[row[0].row].height = 62
    for cell in row:
        cell.font = Font(name='Arial', size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='center')
for cell in ws[2]:
    cell.fill = PatternFill('solid', fgColor='60452F')
    cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
ws.sheet_view.showGridLines = False
wb.save(target)
assert load_workbook(target).active['A3'].value == rows[2][0]
print(target)
