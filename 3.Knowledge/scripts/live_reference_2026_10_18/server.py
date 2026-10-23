"""Local, read-only Excel bridge. No database, writes, external URLs or arbitrary paths."""
import datetime
import hashlib
import io
import json
import math
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
SELECTION = ROOT / '3.Knowledge/reports/account_reference_2026_10_18/selection.json'
ALLOWED = {'SIMPLY JUUN.xlsx', 'Assesment Vape IC 100.xlsx'}

def scalar(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value if value is None or isinstance(value, (str, int, float, bool)) else str(value)

def field(label, coordinate, formulas, cached):
    source, value = formulas[coordinate], cached[coordinate]
    return {'label': label, 'cell': coordinate, 'value': scalar(value.value),
            'formula': source.value if source.data_type == 'f' else None,
            'numberFormat': source.number_format, 'cellType': value.data_type}

def section_rows(title, header, row, formulas, cached):
    fields = [field(str(formulas.cell(header, col).value or openpyxl.utils.get_column_letter(col)), f'{openpyxl.utils.get_column_letter(col)}{row}', formulas, cached)
              for col in range(17, 30) if formulas.cell(header, col).value is not None]
    return {'title': title, 'fields': fields}

def read_accounts(root=ROOT, selections=None):
    selections = selections if selections is not None else json.loads(SELECTION.read_text())['selection']
    profiles = json.loads((root / '3.Knowledge/reports/workbook_reference_2026_10_18/structure_profile.json').read_text())
    loaded, accounts = {}, []
    try:
        for item in selections:
            file = item['sourceFile']
            if file not in ALLOWED:
                raise ValueError('Workbook tidak diizinkan.')
            if file not in loaded:
                source = root / '4. Refrence' / file
                payload = source.read_bytes()
                loaded[file] = (openpyxl.load_workbook(io.BytesIO(payload), data_only=False), openpyxl.load_workbook(io.BytesIO(payload), data_only=True), hashlib.sha256(payload).hexdigest())
            formulas_book, cached_book, digest = loaded[file]
            formulas, cached = formulas_book[item['sheet']], cached_book[item['sheet']]
            profile = next(s for book in profiles if book['file'] == file for s in book['sheets'] if s['sheet'] == item['sheet'])
            row, table = item['sourceRow'], profile['tables'][0]
            name = str(cached.cell(row, 17).value or '').strip()
            if not name:
                raise ValueError('Nama pada baris terpilih kosong; tinjau ulang pemilihan akun.')
            sections = [section_rows('Roster talent', table['header_row'], row, formulas, cached)]
            for extra in profile['tables'][1:]:
                matches = [r for r in range(extra['header_row'] + 1, extra['end_row'] + 1) if str(cached.cell(r, 17).value or '').strip().casefold() == name.casefold()]
                if len(matches) == 1:
                    sections.append(section_rows(extra['heading'], extra['header_row'], matches[0], formulas, cached))
                elif matches:
                    sections.append({'title': extra['heading'], 'fields': [], 'note': 'Nama cocok lebih dari satu baris; tidak digabung otomatis.'})
            overview = []
            for r in range(5, table['header_row'] - 1):
                label = str(cached.cell(r, 18).value or f'Baris {r}')
                for col in range(19, 22):
                    if formulas.cell(r, col).value is not None:
                        overview.append(field(label, f'{openpyxl.utils.get_column_letter(col)}{r}', formulas, cached))
            sections.append({'title': 'Brief, budget dan timeline campaign (bukan nominal pribadi otomatis)', 'fields': overview})
            accounts.append({'id': item['id'], 'name': name, 'sourceFile': file, 'sheet': item['sheet'], 'sourceRow': row, 'sha256': digest, 'sections': sections})
        return {'accounts': accounts, 'readAt': datetime.datetime.now(datetime.timezone.utc).isoformat(), 'mode': 'live_read_only'}
    finally:
        for formulas, cached, _ in loaded.values():
            formulas.close()
            cached.close()

class Handler(BaseHTTPRequestHandler):
    def log_message(self, *_):
        pass  # Never log personal cell values or request data.

    def do_GET(self):
        if self.headers.get('Host') not in ['127.0.0.1:48176', 'localhost:48176']:
            self.send_error(403)
            return
        origin = self.headers.get('Origin')
        if origin and origin not in ['http://127.0.0.1:48175', 'http://127.0.0.1:48174']:
            self.send_error(403)
            return
        if self.path != '/accounts':
            self.send_error(404)
            return
        try:
            data, status = read_accounts(), 200
        except Exception:
            data, status = {'error': 'Excel tidak dapat dibaca. Pastikan file tersedia, nama sheet dan baris pilihan masih valid; lalu muat ulang.'}, 503
        payload = json.dumps(data, ensure_ascii=False, allow_nan=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Cache-Control', 'no-store')
        self.send_header('X-Content-Type-Options', 'nosniff')
        self.end_headers()
        self.wfile.write(payload)

if __name__ == '__main__':
    print('[EXCEL-REFERENCE][INFO] : Read-only loopback bridge on 48176', flush=True)
    ThreadingHTTPServer(('127.0.0.1', 48176), Handler).serve_forever()
