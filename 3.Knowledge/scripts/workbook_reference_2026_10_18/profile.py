"""Read-only structure/profile extraction. Never persists talent/contact/link values."""
import collections
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import urlparse
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / '3.Knowledge/reports/workbook_reference_2026_10_18'
OUT.mkdir(parents=True, exist_ok=True)
FILES = ['SIMPLY JUUN.xlsx', 'Assesment Vape IC 100.xlsx']
result = []
for filename in FILES:
    source = ROOT / '4. Refrence' / filename
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=False)
    cached = openpyxl.load_workbook(source, read_only=True, data_only=True)
    book = dict(file=filename, bytes=source.stat().st_size, sha256=digest, sheets=[])
    for sheet in workbook:
        rows = list(sheet.iter_rows())
        cells = {(c.row, c.column): c for row in rows for c in row if c.value is not None}
        values = {(c.row, c.column): c.value for row in cached[sheet.title].iter_rows() for c in row if c.value is not None}
        headings = [(r, str(c.value).strip()) for (r, col), c in cells.items() if col == 17 and any(str(c.value).strip().startswith(prefix) for prefix in ['TALENT TERPILIH', 'FINAL PROGRESS CREATOR', 'ENGAGEMENT DASBOARD'])]
        headings.sort()
        last_row = max(r for r, col in cells)
        profile = dict(sheet=sheet.title, extent=dict(rows=last_row, columns=max(col for r, col in cells)), nonempty_cells=len(cells), tables=[], metadata_labels=[], formulas=[], cached_errors=[], engagement_checks=dict(checked=0, missing_denominator=[], cached_mismatch=[]))
        for (r, col), cell in cells.items():
            if col == 18 and r < headings[0][0] and isinstance(cell.value, str):
                profile['metadata_labels'].append(dict(cell=cell.coordinate, label=cell.value.strip(), value_cell=f'S{r}'))
            if cell.data_type == 'f':
                profile['formulas'].append(dict(cell=cell.coordinate, formula=cell.value))
                value = values.get((r, col))
                if isinstance(value, str) and value.startswith('#'):
                    profile['cached_errors'].append(dict(cell=cell.coordinate, error=value))
                if re.fullmatch(r'=SUM\(V\d+:Z\d+\)/AA\d+\*100', str(cell.value)):
                    denominator = values.get((r, 27))
                    if not isinstance(denominator, (float, int)) or denominator <= 0:
                        profile['engagement_checks']['missing_denominator'].append(f'AA{r}')
                    else:
                        numerator = sum(values.get((r, col), 0) for col in range(22, 27) if isinstance(values.get((r, col), 0), (float, int)))
                        profile['engagement_checks']['checked'] += 1
                        if not isinstance(value, (float, int)) or abs(numerator / denominator * 100 - value) > 1e-6:
                            profile['engagement_checks']['cached_mismatch'].append(cell.coordinate)
        for index, (heading_row, heading) in enumerate(headings):
            header = heading_row + 1
            end = headings[index + 1][0] - 1 if index + 1 < len(headings) else last_row
            record_rows = [r for r in range(header + 1, end + 1) if (r, 17) in cells and str(cells[(r, 17)].value).strip()]
            table = dict(heading=heading, heading_cell=f'Q{heading_row}', header_row=header, end_row=end, record_count=len(record_rows), columns=[])
            for (r, col), cell in cells.items():
                if r != header or col < 17:
                    continue
                label = str(cell.value).strip()
                counts = collections.Counter()
                invalid = []
                non_numeric = []
                for record in record_rows:
                    value = values.get((record, col))
                    category = 'blank' if value is None or str(value).strip() == '' else 'number' if isinstance(value, (int, float)) else 'text'
                    counts[category] += 1
                    if 'FOLLOWERS' in label.upper() and category == 'text' and not re.fullmatch(r'[\d.,\s]+', str(value)):
                        non_numeric.append(f'{cell.column_letter}{record}')
                    if 'LINK' in label.upper() and category != 'blank':
                        parsed = urlparse(str(value).strip())
                        if parsed.scheme not in ['http', 'https'] or not parsed.netloc or ' ' in parsed.netloc:
                            invalid.append(f'{cell.column_letter}{record}')
                table['columns'].append(dict(cell=cell.coordinate, label=label, types=dict(counts), invalid_url_cells=invalid, non_numeric_follower_cells=non_numeric))
            profile['tables'].append(table)
        book['sheets'].append(profile)
    workbook.close()
    cached.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest, 'Source changed during read'
    book['source_unchanged_after_read'] = True
    result.append(book)
(OUT / 'structure_profile.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
for book in result:
    print(book['file'], [(s['sheet'], [(t['heading_cell'], t['record_count']) for t in s['tables']], len(s['cached_errors'])) for s in book['sheets']])
